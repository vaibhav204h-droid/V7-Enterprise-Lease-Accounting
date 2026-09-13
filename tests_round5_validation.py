from pathlib import Path
from openpyxl import load_workbook
import zipfile

REPORT = Path(__file__).with_name('reports') / 'Lease_Accounting_V7_3_VAIBHAV_SHIPYARD_20260908_ROUND5_FIXED.xlsx'

def test_xlsx_integrity():
    with zipfile.ZipFile(REPORT) as z:
        assert z.testzip() is None

def test_no_cached_formula_blanks_or_errors():
    vals = load_workbook(REPORT, data_only=True)
    formulas = load_workbook(REPORT, data_only=False)
    errors = ('#REF!', '#VALUE!', '#NAME?', '#DIV/0!', '#N/A', '#NUM!', '#NULL!')
    bad = []
    for wsf, wsv in zip(formulas.worksheets, vals.worksheets):
        for row in wsf.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    v = wsv[c.coordinate].value
                    if v in (None, '') or (isinstance(v, str) and v.startswith(errors)):
                        bad.append((wsf.title, c.coordinate, v))
    assert not bad, bad[:20]

def test_escalation_is_not_blank():
    vals = load_workbook(REPORT, data_only=True)
    ws = vals['02_Payment_Schedule']
    for r in range(6, ws.max_row + 1):
        if ws.cell(r, 1).value:
            assert ws.cell(r, 5).value not in (None, '')

def test_control_je_pass():
    vals = load_workbook(REPORT, data_only=True)
    assert vals['00_Control']['B15'].value == 'PASS'
